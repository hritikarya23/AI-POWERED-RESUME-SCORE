"""Unit tests for document text extraction and cleaning."""

import io
import pytest
from pypdf import PdfWriter
import docx

from app.services.extractor import (
    clean_extracted_text,
    extract_from_docx,
    extract_from_pdf,
    extract_from_txt,
    extract_text_from_file,
    ExtractionError,
)


def test_clean_extracted_text():
    raw = "Software  Engineer\r\n\r\n\r\n\r\nwith  experi-\nence in Python\xa0and cloud\x0csystems."
    cleaned = clean_extracted_text(raw)
    assert "experience" in cleaned
    assert "Python and cloud" in cleaned
    assert "\r" not in cleaned
    assert "\x0c" not in cleaned


def test_extract_from_txt():
    content = "John Doe\nEmail: john@example.com\nSkills: Python, FastAPI\nExperience: 5 years building APIs."
    raw_bytes = content.encode("utf-8")
    result = extract_from_txt(raw_bytes)
    assert result["word_count"] > 10
    assert "john@example.com" in result["text"]


def test_extract_from_docx():
    doc = docx.Document()
    doc.add_heading("Jane Smith", level=1)
    doc.add_paragraph("Full Stack Developer with 4 years experience in React and Node.js.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "Docker, AWS"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = extract_from_docx(docx_bytes)
    assert "Jane Smith" in result["text"]
    assert "React" in result["text"]
    assert "Docker" in result["text"]
    assert result["word_count"] > 5


def test_extract_from_pdf():
    # 1. Test empty blank page error handling
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    with pytest.raises(ExtractionError):
        extract_from_pdf(pdf_bytes)

    # 2. Test extraction from real generated PDF resume
    sample_pdf_path = "sample_resumes/alex_rivera_backend_engineer.pdf"
    with open(sample_pdf_path, "rb") as f:
        real_pdf_bytes = f.read()

    res = extract_from_pdf(real_pdf_bytes)
    assert res["page_count"] >= 1
    assert "ALEX RIVERA" in res["text"]
    assert "FastAPI" in res["text"]
    assert "Kubernetes" in res["text"]
    assert res["word_count"] > 100


def test_extract_text_from_file_dispatch():
    content = "Senior Cloud Architect with expertise in Kubernetes and Terraform."
    raw_bytes = content.encode("utf-8")

    result = extract_text_from_file(raw_bytes, "resume.txt")
    assert result["filename"] == "resume.txt"
    assert "Kubernetes" in result["text"]

    with pytest.raises(ExtractionError):
        extract_text_from_file(b"", "empty.txt")


def test_extract_from_image():
    from PIL import Image, ImageDraw
    # Create a clear, high-contrast synthetic image with text
    img = Image.new("RGB", (800, 250), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((30, 40), "Software Developer with Python and Docker skills.", fill=(0, 0, 0))
    draw.text((30, 90), "Experience in building APIs and cloud applications.", fill=(0, 0, 0))
    
    stream = io.BytesIO()
    img.save(stream, format="JPEG")
    img_bytes = stream.getvalue()

    result = extract_text_from_file(img_bytes, "candidate_resume.jpg")
    assert result["filename"] == "candidate_resume.jpg"
    assert result["word_count"] > 3
    # Verify keywords detected by OCR
    assert any(term in result["text"].lower() for term in ["software", "python", "docker", "developer", "experience"])

