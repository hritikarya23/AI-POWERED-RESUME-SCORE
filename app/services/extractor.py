"""Document text extraction service supporting PDF, DOCX, and TXT files."""

import io
import re
from typing import Any, Dict
from pypdf import PdfReader
import docx


class ExtractionError(Exception):
    """Custom exception raised when document extraction fails."""
    pass


def clean_extracted_text(text: str) -> str:
    """Normalize extracted document text for downstream NLP processing."""
    if not text:
        return ""

    # Replace non-breaking spaces and irregular whitespace characters
    text = text.replace("\xa0", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x0c", "\n")  # Form feed characters from PDF page breaks

    # Replace smart/curly quotes and dashes
    text = text.replace("“", '"').replace("”", '"').replace("’", "'").replace("‘", "'")
    text = text.replace("—", " - ").replace("–", " - ")

    # Rejoin hyphenated words split across line breaks: e.g. "experi-\nence" -> "experience"
    text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)

    # Replace 3 or more consecutive newlines with 2 newlines
    text = re.sub(r'\n{3,}', '\n\n', text)

    # Normalize horizontal whitespace on each line
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n')]
    cleaned = '\n'.join(lines)

    return cleaned.strip()


def extract_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """Extract text and metadata from a PDF file."""
    try:
        stream = io.BytesIO(file_bytes)
        reader = PdfReader(stream)

        if reader.is_encrypted:
            try:
                # Attempt empty password decryption
                reader.decrypt("")
            except Exception as e:
                raise ExtractionError(f"PDF is password-protected and cannot be read: {e}")

        page_count = len(reader.pages)
        if page_count == 0:
            raise ExtractionError("PDF file contains no pages.")

        page_texts = []
        for i, page in enumerate(reader.pages):
            try:
                # Try layout mode if available for cleaner column extraction
                text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
            except Exception:
                text = page.extract_text() or ""
            page_texts.append(text)

        full_raw_text = "\n\n".join(page_texts)
        cleaned_text = clean_extracted_text(full_raw_text)

        if not cleaned_text:
            raise ExtractionError(
                "Could not extract any readable text from this PDF. "
                "The PDF might be a scanned image or empty. Please ensure it contains selectable text."
            )

        words = cleaned_text.split()
        return {
            "text": cleaned_text,
            "page_count": page_count,
            "word_count": len(words),
            "char_count": len(cleaned_text),
        }
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to process PDF document: {str(e)}")


def extract_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """Extract text from a Word DOCX file, including paragraphs and tables."""
    try:
        stream = io.BytesIO(file_bytes)
        doc = docx.Document(stream)

        extracted_lines = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                extracted_lines.append(para.text.strip())

        # Also extract text from tables (often used in resumes for headers/skills)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    extracted_lines.append(" | ".join(row_cells))

        full_text = "\n".join(extracted_lines)
        cleaned_text = clean_extracted_text(full_text)

        if not cleaned_text:
            raise ExtractionError("The DOCX document is empty or contains no readable text.")

        words = cleaned_text.split()
        # Estimate page count (~400 words per page for resumes)
        estimated_pages = max(1, (len(words) + 399) // 400)

        return {
            "text": cleaned_text,
            "page_count": estimated_pages,
            "word_count": len(words),
            "char_count": len(cleaned_text),
        }
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to process Word DOCX document: {str(e)}")


def extract_from_txt(file_bytes: bytes) -> Dict[str, Any]:
    """Extract and decode text from a plain text file."""
    try:
        # Try UTF-8 first, fallback to Latin-1
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")

        cleaned_text = clean_extracted_text(text)
        if not cleaned_text:
            raise ExtractionError("The text file is empty.")

        words = cleaned_text.split()
        estimated_pages = max(1, (len(words) + 399) // 400)

        return {
            "text": cleaned_text,
            "page_count": estimated_pages,
            "word_count": len(words),
            "char_count": len(cleaned_text),
        }
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to process text file: {str(e)}")


def extract_from_image(file_bytes: bytes, filename: str = "resume.jpg") -> Dict[str, Any]:
    """Extract text from an image (JPG, JPEG, PNG, WEBP, BMP) using Groq Vision or local fallback."""
    # First priority: Fast & accurate Groq Vision
    try:
        from app.services.groq_service import is_groq_available, extract_text_from_image_groq
        if is_groq_available():
            text = extract_text_from_image_groq(file_bytes, filename)
            if text and len(text.strip()) > 10:
                cleaned = clean_extracted_text(text)
                words = cleaned.split()
                return {
                    "text": cleaned,
                    "page_count": 1,
                    "word_count": len(words),
                    "char_count": len(cleaned),
                }
    except Exception:
        pass

    # Second priority: Local OCR if easyocr and numpy are available
    try:
        import numpy as np
        import easyocr
        from PIL import Image

        global _OCR_READER
        if "_OCR_READER" not in globals() or _OCR_READER is None:
            _OCR_READER = easyocr.Reader(["en"], gpu=False, verbose=False)

        image = Image.open(io.BytesIO(file_bytes))
        if image.mode != "RGB":
            image = image.convert("RGB")

        extracted_chunks = _OCR_READER.readtext(np.array(image), detail=0, paragraph=True)
        raw_text = "\n\n".join(extracted_chunks)
        cleaned_text = clean_extracted_text(raw_text)

        if not cleaned_text or len(cleaned_text.strip()) < 10:
            raise ExtractionError(
                "Could not detect clear readable text in this image. "
                "Please make sure the image is sharp, well-lit, and contains readable resume text."
            )

        words = cleaned_text.split()
        return {
            "text": cleaned_text,
            "page_count": 1,
            "word_count": len(words),
            "char_count": len(cleaned_text),
        }
    except (ImportError, ModuleNotFoundError):
        raise ExtractionError(
            "Could not extract text from this image. "
            "Please upload your resume in PDF (.pdf), Word (.docx), or plain text (.txt) format, "
            "or copy and paste your resume text into the text area."
        )
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to process resume image: {str(e)}")


IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff")


def extract_text_from_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Main entry point for document extraction.
    Automatically detects format from filename (PDF, DOCX, TXT, JPG, PNG, WEBP)
    and returns structured text + statistics.
    """
    if not file_bytes:
        raise ExtractionError("Uploaded file is empty (0 bytes).")

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        data = extract_from_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        data = extract_from_docx(file_bytes)
    elif lower_name.endswith(".txt") or lower_name.endswith(".md"):
        data = extract_from_txt(file_bytes)
    elif any(lower_name.endswith(ext) for ext in IMAGE_EXTENSIONS):
        data = extract_from_image(file_bytes, filename)
    else:
        # Fallback: attempt PDF, then Image, then text
        try:
            data = extract_from_pdf(file_bytes)
        except Exception:
            try:
                data = extract_from_image(file_bytes, filename)
            except Exception:
                try:
                    data = extract_from_txt(file_bytes)
                except Exception:
                    raise ExtractionError(
                        f"Unsupported file format '{filename}'. Please upload a PDF, Word DOCX, Text, or Image file (JPG/PNG)."
                    )

    data["filename"] = filename
    return data

